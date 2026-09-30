import os
import re
import uuid
import shutil
import time
import polars as pl
from typing import Tuple, List, Dict, Any, Optional
from app.models import CleaningOptions

def load_dataset(file_path: str) -> pl.DataFrame:
    """Loads CSV, XLSX, TSV, Parquet, JSON files with multi-engine resilience (Polars + DuckDB)."""
    ext = os.path.splitext(file_path)[1].lower()

    if ext in [".csv", ".txt", ".tsv"]:
        sep = "\t" if ext == ".tsv" else ","
        # 1. Try standard Polars reader
        try:
            return pl.read_csv(
                file_path,
                separator=sep,
                infer_schema_length=10000,
                ignore_errors=True,
                truncate_ragged_lines=True
            )
        except Exception:
            pass

        # 2. Try DuckDB reader (resilient to unescaped quotes, ragged lines, irregular escaping)
        try:
            import duckdb
            rel = duckdb.read_csv(file_path, ignore_errors=True, auto_detect=True)
            return rel.pl()
        except Exception:
            pass

        # 3. Try Polars with other common separators
        for alt_sep in [";", "\t", "|", ","]:
            try:
                return pl.read_csv(
                    file_path,
                    separator=alt_sep,
                    infer_schema_length=10000,
                    ignore_errors=True,
                    truncate_ragged_lines=True
                )
            except Exception:
                continue

        # 4. Try DuckDB with all_varchar=True
        try:
            import duckdb
            rel = duckdb.read_csv(file_path, ignore_errors=True, all_varchar=True, auto_detect=True)
            return rel.pl()
        except Exception as e:
            raise ValueError(f"Could not parse delimited file {file_path}: {str(e)}")

    elif ext in [".xlsx", ".xls"]:
        try:
            return pl.read_excel(file_path)
        except Exception:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                return pl.DataFrame()
            headers = [str(h) if h is not None else f"col_{i}" for i, h in enumerate(rows[0])]
            data = rows[1:]
            return pl.DataFrame(data, schema=headers, orient="row")

    elif ext == ".parquet":
        return pl.read_parquet(file_path)

    elif ext in [".json", ".jsonl", ".ndjson"]:
        try:
            return pl.read_json(file_path)
        except Exception:
            return pl.read_ndjson(file_path)

    else:
        # Fallback to DuckDB or Polars
        try:
            import duckdb
            rel = duckdb.read_csv(file_path, ignore_errors=True, auto_detect=True)
            return rel.pl()
        except Exception:
            return pl.read_csv(
                file_path,
                infer_schema_length=10000,
                ignore_errors=True,
                truncate_ragged_lines=True
            )

def get_empty_column_names(df: pl.DataFrame, threshold_pct: Optional[float] = None) -> List[str]:
    """Find columns that are 100% blank (or exceed null/blank threshold percentage)."""
    empty_cols = []
    total_rows = df.height
    if total_rows == 0:
        return []

    for col in df.columns:
        series = df[col]
        null_count = series.null_count()
        
        # If it's a string column, also count empty or all-whitespace strings
        if series.dtype == pl.Utf8 or series.dtype == pl.String:
            # Count nulls + blank strings
            blank_mask = series.is_null() | (series.str.strip_chars() == "")
            blank_count = blank_mask.sum()
        else:
            blank_count = null_count

        pct_blank = (blank_count / total_rows) * 100.0
        
        if threshold_pct is not None:
            if pct_blank >= threshold_pct:
                empty_cols.append(col)
        else:
            if blank_count == total_rows:
                empty_cols.append(col)

    return empty_cols

def apply_cleaning_pipeline(df: pl.DataFrame, options: CleaningOptions) -> pl.DataFrame:
    """Applies all cleaning, extraction, filtering, and sorting steps."""
    
    rename_map = options.rename_columns or {}

    # 0. Column Selection & Renaming (Executed FIRST so all downstream operations target selected & renamed columns)
    if options.selected_columns:
        valid_cols = [c for c in options.selected_columns if c in df.columns]
        if valid_cols:
            df = df.select(valid_cols)

    if options.rename_columns:
        valid_renames = {k: v for k, v in options.rename_columns.items() if k in df.columns and v and v.strip()}
        if valid_renames:
            df = df.rename(valid_renames)

    # 1. Blank Columns Removal
    if options.drop_col_if_any_null and df.height > 0:
        # Strict: Drop ANY column that has ANY null or blank value (even if only 1 row is blank)
        cols_to_keep = []
        for col in df.columns:
            series = df[col]
            if series.null_count() == 0:
                if series.dtype in (pl.Utf8, pl.String):
                    if (series.str.strip_chars() == "").sum() == 0:
                        cols_to_keep.append(col)
                else:
                    cols_to_keep.append(col)
        if cols_to_keep:
            df = df.select(cols_to_keep)
    elif options.remove_all_blank_columns or (options.blank_column_threshold_pct is not None):
        threshold = options.blank_column_threshold_pct if options.blank_column_threshold_pct is not None else 100.0
        cols_to_drop = get_empty_column_names(df, threshold_pct=threshold)
        if cols_to_drop:
            remaining = [c for c in df.columns if c not in cols_to_drop]
            if remaining:
                df = df.select(remaining)

    # 2. Strict Blank Rows Removal (Drop row if ANY cell is null/blank)
    if options.drop_row_if_any_null and df.height > 0 and len(df.columns) > 0:
        is_blank_exprs = [
            pl.col(c).is_null() | (pl.col(c).cast(pl.Utf8, strict=False).str.strip_chars() == "")
            for c in df.columns
        ]
        df = df.filter(~pl.any_horizontal(is_blank_exprs))

    # 3. Drop rows with nulls in specific required column headers
    elif options.drop_rows_with_null_in_cols and df.height > 0:
        null_target_cols = []
        for c in options.drop_rows_with_null_in_cols:
            if c in df.columns:
                null_target_cols.append(c)
            elif c in rename_map and rename_map[c] in df.columns:
                null_target_cols.append(rename_map[c])
        if null_target_cols:
            is_blank_exprs = [
                pl.col(c).is_null() | (pl.col(c).cast(pl.Utf8, strict=False).str.strip_chars() == "")
                for c in null_target_cols
            ]
            if (options.drop_null_mode or "any").lower() == "all":
                df = df.filter(~pl.all_horizontal(is_blank_exprs))
            else:
                # Any blank in selected heads drops the row
                df = df.filter(~pl.any_horizontal(is_blank_exprs))

    # 4. Remove 100% empty rows (if strict row drop wasn't already active)
    elif options.remove_all_blank_rows and df.height > 0 and len(df.columns) > 0:
        conditions = []
        for col in df.columns:
            if df[col].dtype in (pl.Utf8, pl.String):
                conditions.append(pl.col(col).is_null() | (pl.col(col).str.strip_chars() == ""))
            else:
                conditions.append(pl.col(col).is_null())
        df = df.filter(~pl.all_horizontal(conditions))

    # 4. Text & Whitespace Scrubbing
    if options.text_casing or options.trim_whitespace or options.collapse_multiple_spaces:
        casing = options.text_casing.lower() if options.text_casing else None
        
        if options.casing_target_columns:
            target_casing_cols = set()
            for c in options.casing_target_columns:
                if c in df.columns:
                    target_casing_cols.add(c)
                elif c in rename_map and rename_map[c] in df.columns:
                    target_casing_cols.add(rename_map[c])
            if not target_casing_cols:
                target_casing_cols = set(df.columns)
        else:
            target_casing_cols = set(df.columns)

        text_exprs = []
        for col in df.columns:
            is_target_casing = (col in target_casing_cols)
            is_string_dtype = df[col].dtype in (pl.Utf8, pl.String, pl.Categorical, pl.Object)
            
            if is_target_casing or is_string_dtype or options.trim_whitespace or options.collapse_multiple_spaces:
                expr = pl.col(col).cast(pl.Utf8, strict=False)
                if options.trim_whitespace:
                    expr = expr.str.strip_chars()
                if options.collapse_multiple_spaces:
                    expr = expr.str.replace_all(r"[ \t]+", " ")
                if casing and is_target_casing:
                    if casing in ("uppercase", "upper", "all_caps", "full_case"):
                        expr = expr.str.to_uppercase()
                    elif casing in ("lowercase", "lower"):
                        expr = expr.str.to_lowercase()
                    elif casing in ("titlecase", "title", "capitalize_words", "1st_letters"):
                        expr = expr.str.to_titlecase()
                    elif casing in ("capitalize", "sentencecase", "sentence_case", "1st_letter_only"):
                        expr = expr.str.slice(0, 1).str.to_uppercase() + expr.str.slice(1).str.to_lowercase()
                text_exprs.append(expr.alias(col))
                
        if text_exprs:
            df = df.with_columns(text_exprs)

    # 5. Clean Numeric Symbols (strip $, ₹, €, commas, %)
    if options.clean_numeric_symbols:
        if options.numeric_target_columns:
            target_cols = [c for c in options.numeric_target_columns if c in df.columns]
        else:
            # Auto-detect candidate columns that actually contain currencies or numbers
            string_cols = [c for c in df.columns if df[c].dtype in (pl.Utf8, pl.String, pl.Categorical, pl.Object)]
            target_cols = []
            for c in string_cols:
                # Check column name heuristics or sample values
                is_name_hint = bool(re.search(r"(?i)(amount|price|cost|salary|fee|total|paid|rate|balance|revenue|tax|payment)", c))
                sample_vals = df[c].drop_nulls().head(20).to_list()
                if not sample_vals:
                    continue
                valid_num_matches = 0
                for v in sample_vals:
                    v_str = str(v).strip()
                    cleaned_sample = re.sub(r"[\$,₹,€,£,%\s]", "", v_str)
                    try:
                        float(cleaned_sample)
                        valid_num_matches += 1
                    except ValueError:
                        pass
                # If more than 60% of sampled values look numeric or name strongly hints
                if valid_num_matches / len(sample_vals) >= 0.6 or (is_name_hint and valid_num_matches > 0):
                    target_cols.append(c)

        num_exprs = []
        for col in target_cols:
            if col in df.columns and df[col].dtype in (pl.Utf8, pl.String):
                cleaned = (
                    pl.col(col)
                    .str.replace_all(r"[\$,₹,€,£,%\s]", "")
                    .cast(pl.Float64, strict=False)
                    .alias(col)
                )
                num_exprs.append(cleaned)
        if num_exprs:
            df = df.with_columns(num_exprs)

    # 6. Bulk Extraction / Match Filter (Crucial: User gave long list of IDs/values)
    if options.bulk_match_column and options.bulk_match_values and options.bulk_match_column in df.columns:
        col = options.bulk_match_column
        raw_values = [v.strip() for v in options.bulk_match_values if v and v.strip()]
        
        if raw_values:
            match_type = (options.bulk_match_type or "exact").lower()
            mode = (options.bulk_match_mode or "keep").lower()

            if match_type == "case_insensitive":
                lower_vals = set(v.lower() for v in raw_values)
                match_condition = pl.col(col).cast(pl.Utf8).str.to_lowercase().is_in(list(lower_vals))
            elif match_type == "contains":
                # For contains, construct regex if reasonable size or batch
                safe_patterns = [re.escape(v) for v in raw_values[:500]]
                combined_regex = "|".join(safe_patterns)
                match_condition = pl.col(col).cast(pl.Utf8).str.contains(combined_regex)
            else:  # exact
                val_set = set(raw_values)
                match_condition = pl.col(col).cast(pl.Utf8).is_in(list(val_set))

            if mode == "keep":
                df = df.filter(match_condition)
            else:  # exclude
                df = df.filter(~match_condition)

    # 7. Conditional Filter Rules
    if options.filter_rules and df.height > 0:
        for rule in options.filter_rules:
            r_col = rule.get("column")
            r_op = rule.get("op")
            r_val = rule.get("val")
            if not r_col or r_col not in df.columns or not r_op:
                continue

            expr_col = pl.col(r_col)
            col_type = df[r_col].dtype

            if r_op == "is_null":
                df = df.filter(expr_col.is_null())
            elif r_op == "is_not_null":
                df = df.filter(expr_col.is_not_null())
            elif r_op == "=":
                if col_type in (pl.Int64, pl.Int32, pl.Float64, pl.Float32):
                    try:
                        df = df.filter(expr_col == float(r_val))
                    except ValueError:
                        pass
                else:
                    df = df.filter(expr_col.cast(pl.Utf8) == str(r_val))
            elif r_op == "!=":
                if col_type in (pl.Int64, pl.Int32, pl.Float64, pl.Float32):
                    try:
                        df = df.filter(expr_col != float(r_val))
                    except ValueError:
                        pass
                else:
                    df = df.filter(expr_col.cast(pl.Utf8) != str(r_val))
            elif r_op in (">", ">=", "<", "<="):
                try:
                    num_val = float(r_val)
                    if r_op == ">":
                        df = df.filter(expr_col > num_val)
                    elif r_op == ">=":
                        df = df.filter(expr_col >= num_val)
                    elif r_op == "<":
                        df = df.filter(expr_col < num_val)
                    elif r_op == "<=":
                        df = df.filter(expr_col <= num_val)
                except ValueError:
                    pass
            elif r_op == "contains":
                df = df.filter(expr_col.cast(pl.Utf8).str.contains(re.escape(str(r_val))))
            elif r_op == "starts_with":
                df = df.filter(expr_col.cast(pl.Utf8).str.starts_with(str(r_val)))
            elif r_op == "ends_with":
                df = df.filter(expr_col.cast(pl.Utf8).str.ends_with(str(r_val)))

    # 8. Deduplication (Eliminates Repeating People and Names)
    if options.remove_duplicates and df.height > 0:
        mode = getattr(options, "dedupe_names_mode", "unique_person") or "unique_person"

        # Mode A: Strict Unique First Name (Never repeat any first name in the file)
        if mode == "unique_first_name":
            first_col = next((c for c in df.columns if re.search(r'(?i)^(first[_\s]?name|first)$', c)), None)
            if not first_col:
                first_col = next((c for c in df.columns if re.search(r'(?i)(first[_\s]?name)', c)), None)
            if not first_col:
                first_col = next((c for c in df.columns if 'name' in c.lower()), None)
            subset = [first_col] if first_col else None

        # Mode B: Unique Full People (First Name + Last Name or Full Name)
        elif mode == "unique_person":
            first_col = next((c for c in df.columns if re.search(r'(?i)^(first[_\s]?name|first)$', c)), None)
            last_col = next((c for c in df.columns if re.search(r'(?i)^(last[_\s]?name|last)$', c)), None)
            full_col = next((c for c in df.columns if re.search(r'(?i)^(full[_\s]?name|full_name)$', c)), None)

            if first_col and last_col:
                subset = [first_col, last_col]
            elif full_col:
                subset = [full_col]
            elif first_col:
                subset = [first_col]
            else:
                subset = [c for c in (options.duplicate_columns or []) if c in df.columns]
                if not subset:
                    name_cols = [c for c in df.columns if 'name' in c.lower()]
                    subset = name_cols if name_cols else None

        # Mode C: Unique Full Name
        elif mode == "unique_full_name":
            full_col = next((c for c in df.columns if re.search(r'(?i)^(full[_\s]?name|full_name|name)$', c)), None)
            if full_col:
                subset = [full_col]
            else:
                subset = [c for c in df.columns if 'name' in c.lower()]

        # Mode D: Strict All Columns
        elif mode == "all_columns":
            subset = None

        # Mode E: Custom Columns / Fallback
        else:
            subset = [c for c in (options.duplicate_columns or []) if c in df.columns]
            if not subset:
                # Intelligent fallback: If user enabled dedupe but subset is empty,
                # default to Person / Name columns rather than requiring all 60+ columns to be identical!
                first_col = next((c for c in df.columns if re.search(r'(?i)^(first[_\s]?name|first)$', c)), None)
                last_col = next((c for c in df.columns if re.search(r'(?i)^(last[_\s]?name|last)$', c)), None)
                full_col = next((c for c in df.columns if re.search(r'(?i)^(full[_\s]?name|full_name|name)$', c)), None)
                if first_col and last_col:
                    subset = [first_col, last_col]
                elif full_col:
                    subset = [full_col]
                else:
                    name_cols = [c for c in df.columns if 'name' in c.lower()]
                    subset = name_cols if name_cols else None

        keep_mode = "last" if options.duplicate_keep == "last" else "first"

        # Trim & normalize string columns in temporary subset for robust matching without mutating original column text
        if subset:
            temp_subset = []
            temp_exprs = []
            for sc in subset:
                if sc in df.columns:
                    temp_col = f"__dedupe_norm_{sc}"
                    temp_subset.append(temp_col)
                    if df[sc].dtype in (pl.Utf8, pl.String):
                        temp_exprs.append(pl.col(sc).str.strip_chars().str.to_lowercase().alias(temp_col))
                    else:
                        temp_exprs.append(pl.col(sc).alias(temp_col))
            if temp_exprs:
                df = df.with_columns(temp_exprs)
                df = df.unique(subset=temp_subset, keep=keep_mode, maintain_order=True)
                df = df.drop(temp_subset)
            else:
                df = df.unique(subset=subset, keep=keep_mode, maintain_order=True)
        else:
            df = df.unique(subset=None, keep=keep_mode, maintain_order=True)

    # 9. Missing Value Imputation
    if options.fill_nulls_mode and df.height > 0:
        target_cols = [c for c in (options.fill_nulls_columns or df.columns) if c in df.columns]
        mode = options.fill_nulls_mode.lower()
        fill_exprs = []

        for col in target_cols:
            col_type = df[col].dtype
            if mode in ("empty_str", "blank"):
                if col_type in (pl.Utf8, pl.String):
                    fill_exprs.append(pl.col(col).fill_null("").alias(col))
                else:
                    fill_exprs.append(pl.col(col).cast(pl.Utf8, strict=False).fill_null("").alias(col))
            elif mode == "custom" and options.fill_nulls_custom_value is not None:
                if col_type in (pl.Utf8, pl.String):
                    fill_exprs.append(pl.col(col).fill_null(options.fill_nulls_custom_value).alias(col))
                else:
                    try:
                        val = float(options.fill_nulls_custom_value)
                        fill_exprs.append(pl.col(col).fill_null(val).alias(col))
                    except ValueError:
                        fill_exprs.append(pl.col(col).cast(pl.Utf8).fill_null(options.fill_nulls_custom_value).alias(col))
            elif mode == "zero":
                if col_type in (pl.Int64, pl.Int32, pl.Float64, pl.Float32):
                    fill_exprs.append(pl.col(col).fill_null(0).alias(col))
                elif col_type in (pl.Utf8, pl.String):
                    fill_exprs.append(pl.col(col).fill_null("0").alias(col))
            elif mode == "forward":
                fill_exprs.append(pl.col(col).forward_fill().alias(col))
            elif mode == "backward":
                fill_exprs.append(pl.col(col).backward_fill().alias(col))
            elif mode == "mean" and col_type in (pl.Int64, pl.Int32, pl.Float64, pl.Float32):
                mean_val = df[col].mean()
                if mean_val is not None:
                    fill_exprs.append(pl.col(col).fill_null(mean_val).alias(col))
            elif mode == "median" and col_type in (pl.Int64, pl.Int32, pl.Float64, pl.Float32):
                med_val = df[col].median()
                if med_val is not None:
                    fill_exprs.append(pl.col(col).fill_null(med_val).alias(col))

        if fill_exprs:
            df = df.with_columns(fill_exprs)

    # 10. Sorting (Only when explicitly specified by the user)
    if options.sort_columns and df.height > 0:
        valid_sort_cols = [c for c in options.sort_columns if c in df.columns]
        if valid_sort_cols:
            descending = options.sort_descending or [False] * len(valid_sort_cols)
            if len(descending) < len(valid_sort_cols):
                descending.extend([False] * (len(valid_sort_cols) - len(descending)))
            df = df.sort(by=valid_sort_cols, descending=descending, nulls_last=options.sort_nulls_last)

    return df

def export_dataset(df: pl.DataFrame, output_path: str, fmt: str = "csv", delimiter: str = ",") -> str:
    """Exports DataFrame to disk as a single unified file (Excel .xlsx, CSV, Parquet, TSV, JSON)."""
    fmt = fmt.lower()
    os.makedirs(os.path.dirname(output_path), exist_ok=True)

    if fmt in ["xlsx", "excel"]:
        # Write to a SINGLE Excel sheet in a single .xlsx file
        max_excel_rows = 1048500
        if df.height <= max_excel_rows:
            df.write_excel(output_path, worksheet="Cleaned_Data")
        else:
            # If strictly beyond Excel row limit (1,048,576 rows)
            import openpyxl
            wb = openpyxl.Workbook()
            wb.remove(wb.active)
            chunk_size = 1000000
            for i in range(0, df.height, chunk_size):
                sheet_num = (i // chunk_size) + 1
                ws = wb.create_sheet(title=f"Cleaned_Data_Part_{sheet_num}")
                chunk = df.slice(i, min(chunk_size, df.height - i))
                ws.append(chunk.columns)
                for row in chunk.iter_rows():
                    ws.append(list(row))
            wb.save(output_path)
    elif fmt == "csv":
        df.write_csv(output_path, separator=delimiter)
    elif fmt == "tsv":
        df.write_csv(output_path, separator="\t")
    elif fmt == "parquet":
        df.write_parquet(output_path, compression="zstd")
    elif fmt == "json":
        df.write_json(output_path)
    elif fmt == "html":
        export_html_table(df, output_path)
    else:
        df.write_excel(output_path, worksheet="Cleaned_Data")

    return output_path

def export_html_table(df: pl.DataFrame, output_path: str, max_embedded_rows: int = 15000) -> str:
    """Exports a Polars DataFrame to a self-contained, responsive, searchable HTML table report."""
    import html as html_lib
    import json
    import math

    cols = df.columns
    total_rows = df.height
    total_cols = df.width

    export_df = df.head(max_embedded_rows) if df.height > max_embedded_rows else df
    headers_json = json.dumps(cols)

    def clean_val(v):
        if v is None:
            return ""
        if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
            return ""
        return str(v)

    rows_data = [[clean_val(cell) for cell in row] for row in export_df.iter_rows()]
    rows_json = json.dumps(rows_data)
    truncated_msg = f"Showing first {max_embedded_rows:,} rows of {total_rows:,} total rows for instant in-browser viewing." if total_rows > max_embedded_rows else f"All {total_rows:,} rows loaded."

    html_content = f"""<!DOCTYPE html>
<html lang="en" class="dark">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Cleaned Data Report • SheetForge Pro</title>
    <script src="https://cdn.tailwindcss.com"></script>
    <script src="https://unpkg.com/lucide@latest"></script>
    <style>
        body {{ background-color: #0b0f19; color: #f9fafb; font-family: system-ui, -apple-system, sans-serif; }}
        .table-scroll {{ max-height: 72vh; overflow: auto; }}
        th {{ position: sticky; top: 0; background: #111827; z-index: 10; }}
        tr:hover td {{ background: rgba(55, 65, 81, 0.4); }}
        ::-webkit-scrollbar {{ width: 6px; height: 6px; }}
        ::-webkit-scrollbar-track {{ background: #0b0f19; }}
        ::-webkit-scrollbar-thumb {{ background: #374151; border-radius: 3px; }}
    </style>
</head>
<body class="p-4 sm:p-8 min-h-screen">
    <div class="max-w-7xl mx-auto space-y-6">
        <!-- Header -->
        <div class="flex flex-col sm:flex-row items-start sm:items-center justify-between gap-4 p-5 rounded-2xl bg-[#111827] border border-gray-800 shadow-xl">
            <div>
                <h1 class="text-xl font-bold text-white flex items-center space-x-2">
                    <span class="w-3 h-3 rounded-full bg-emerald-400"></span>
                    <span>Cleaned Dataset Report</span>
                </h1>
                <p class="text-xs text-gray-400 mt-1">{truncated_msg}</p>
            </div>
            <div class="flex items-center space-x-3 text-xs">
                <span class="px-3 py-1.5 rounded-lg bg-gray-800 text-gray-300 font-mono"><b>{total_rows:,}</b> Rows</span>
                <span class="px-3 py-1.5 rounded-lg bg-gray-800 text-gray-300 font-mono"><b>{total_cols:,}</b> Columns</span>
                <button onclick="downloadCSV()" class="px-4 py-2 rounded-lg bg-emerald-600 hover:bg-emerald-500 font-bold text-white shadow-md transition-all">
                    Export as CSV
                </button>
            </div>
        </div>

        <!-- Controls: Search & Pagination -->
        <div class="flex flex-col sm:flex-row items-center justify-between gap-4">
            <div class="w-full sm:w-80 relative">
                <input type="text" id="searchInput" placeholder="Filter rows in real time..." oninput="handleSearch()"
                    class="w-full pl-9 pr-4 py-2 text-xs rounded-xl bg-gray-900 border border-gray-700 text-white placeholder-gray-500 focus:outline-none focus:border-indigo-500">
                <span class="absolute left-3 top-2 text-gray-400 text-xs">🔍</span>
            </div>
            <div class="flex items-center space-x-2 text-xs text-gray-300">
                <span>Page:</span>
                <button onclick="prevPage()" class="px-2.5 py-1 rounded bg-gray-800 hover:bg-gray-700 text-white">&lt;</button>
                <span id="pageInfo" class="font-mono px-2">1 / 1</span>
                <button onclick="nextPage()" class="px-2.5 py-1 rounded bg-gray-800 hover:bg-gray-700 text-white">&gt;</button>
                <select id="pageSizeSelect" onchange="changePageSize()" class="bg-gray-900 border border-gray-700 rounded px-2 py-1 text-xs text-gray-200">
                    <option value="50">50 / page</option>
                    <option value="100" selected>100 / page</option>
                    <option value="250">250 / page</option>
                    <option value="500">500 / page</option>
                </select>
            </div>
        </div>

        <!-- Interactive Table -->
        <div class="table-scroll rounded-2xl border border-gray-800 shadow-2xl bg-gray-900/80">
            <table class="w-full text-left text-xs whitespace-nowrap divide-y divide-gray-800">
                <thead id="tableHead" class="divide-x divide-gray-800 text-gray-300 font-semibold uppercase tracking-wider text-[11px]"></thead>
                <tbody id="tableBody" class="divide-y divide-gray-800 text-gray-200"></tbody>
            </table>
        </div>
    </div>

    <script>
        const columns = {headers_json};
        const allRows = {rows_json};
        let filteredRows = allRows;
        let currentPage = 1;
        let pageSize = 100;

        function renderHeaders() {{
            const tr = document.createElement('tr');
            columns.forEach(col => {{
                const th = document.createElement('th');
                th.className = 'px-4 py-3 bg-[#111827] border-b border-gray-800';
                th.textContent = col;
                tr.appendChild(th);
            }});
            document.getElementById('tableHead').appendChild(tr);
        }}

        function renderBody() {{
            const tbody = document.getElementById('tableBody');
            tbody.innerHTML = '';
            const start = (currentPage - 1) * pageSize;
            const end = Math.min(start + pageSize, filteredRows.length);
            const slice = filteredRows.slice(start, end);

            const fragment = document.createDocumentFragment();
            slice.forEach(row => {{
                const tr = document.createElement('tr');
                tr.className = 'hover:bg-gray-800/50 transition-colors';
                row.forEach(val => {{
                    const td = document.createElement('td');
                    td.className = 'px-4 py-2 border-b border-gray-800/60 max-w-xs truncate font-mono text-[11px]';
                    td.textContent = val !== null && val !== undefined ? val : '';
                    tr.appendChild(td);
                }});
                fragment.appendChild(tr);
            }});
            tbody.appendChild(fragment);

            const totalPages = Math.ceil(filteredRows.length / pageSize) || 1;
            document.getElementById('pageInfo').textContent = `${{currentPage}} / ${{totalPages}} (${{filteredRows.length.toLocaleString()}} records)`;
        }}

        function handleSearch() {{
            const q = document.getElementById('searchInput').value.toLowerCase().trim();
            if (!q) {{
                filteredRows = allRows;
            }} else {{
                filteredRows = allRows.filter(row => row.some(cell => String(cell).toLowerCase().includes(q)));
            }}
            currentPage = 1;
            renderBody();
        }}

        function nextPage() {{
            const totalPages = Math.ceil(filteredRows.length / pageSize) || 1;
            if (currentPage < totalPages) {{ currentPage++; renderBody(); }}
        }}

        function prevPage() {{
            if (currentPage > 1) {{ currentPage--; renderBody(); }}
        }}

        function changePageSize() {{
            pageSize = parseInt(document.getElementById('pageSizeSelect').value) || 100;
            currentPage = 1;
            renderBody();
        }}

        function downloadCSV() {{
            let csv = columns.map(c => `"${{c.replace(/"/g, '""')}}"`).join(',') + '\\n';
            filteredRows.forEach(r => {{
                csv += r.map(cell => `"${{String(cell || '').replace(/"/g, '""')}}"`).join(',') + '\\n';
            }});
            const blob = new Blob([csv], {{ type: 'text/csv;charset=utf-8;' }});
            const link = document.createElement('a');
            link.href = URL.createObjectURL(blob);
            link.download = 'Cleaned_Data.csv';
            link.click();
        }}

        renderHeaders();
        renderBody();
    </script>
</body>
</html>"""

    with open(output_path, "w", encoding="utf-8") as f:
        f.write(html_content)

    return output_path

def get_column_distinct_values(df: pl.DataFrame, column: str, limit: int = 500) -> Dict[str, Any]:
    """Retrieves top distinct values and counts for a specific column to power category selection."""
    if column not in df.columns:
        return {"column": column, "total_unique": 0, "values": []}

    s = df[column]
    total_unique = s.n_unique()
    vc = s.value_counts(sort=True)

    values = []
    for row in vc.head(limit).iter_rows():
        val = row[0]
        cnt = int(row[1])
        if val is None or (isinstance(val, str) and not val.strip()):
            display_val = "(Blank / Unspecified)"
            raw_val = ""
        else:
            raw_val = str(val).strip()
            display_val = raw_val

        values.append({
            "value": raw_val,
            "display": display_val,
            "count": cnt
        })

    return {
        "column": column,
        "total_unique": total_unique,
        "values": values
    }

def sanitize_sheet_title(name: str, seen: set) -> str:
    """Sanitizes category name to a valid Excel sheet name (max 31 chars, no forbidden chars)."""
    raw = str(name).strip() if name is not None else ""
    if not raw or raw.lower() in ("null", "none", "nan", ""):
        base = "Blank_Unspecified"
    else:
        # Excel forbids \ / ? * [ ] : and quotes at edges
        base = re.sub(r'[\\/*?:\[\]]', '_', raw).strip("'").strip()
        if not base:
            base = "Category"

    # Truncate to 31 chars with collision suffix
    cand = base[:31]
    counter = 1
    while cand.lower() in seen:
        suffix = f"_{counter}"
        cand = base[:31 - len(suffix)] + suffix
        counter += 1

    seen.add(cand.lower())
    return cand

def export_by_category(
    df: pl.DataFrame,
    split_column: str,
    export_mode: str,
    file_format: str,
    output_dir: str,
    base_filename: str,
    selected_categories: Optional[List[str]] = None,
    include_summary_sheet: bool = True,
    secondary_sort_col: Optional[str] = None,
    secondary_sort_desc: bool = False
) -> Dict[str, Any]:
    """
    Sorts and exports dataset divided by country or other category basis.
    Modes:
      - 'sheets': Single Excel workbook (.xlsx) where each country has its own worksheet tab + Overview sheet.
      - 'zip': Downloadable .zip archive containing separate files per country (.xlsx or .csv).
      - 'sorted_unified': Single unified file sorted & grouped primarily by country.
    """
    import zipfile
    import xlsxwriter

    os.makedirs(output_dir, exist_ok=True)
    if split_column not in df.columns:
        raise ValueError(f"Category column '{split_column}' does not exist in dataset.")

    # 1. Secondary sort if specified
    sort_cols = [split_column]
    sort_descs = [False]
    if secondary_sort_col and secondary_sort_col in df.columns and secondary_sort_col != split_column:
        sort_cols.append(secondary_sort_col)
        sort_descs.append(secondary_sort_desc)

    # Sort entire dataset first so records are naturally grouped
    df_sorted = df.sort(by=sort_cols, descending=sort_descs, nulls_last=True)

    # 2. Filter by selected_categories if provided
    if selected_categories is not None and len(selected_categories) > 0:
        sel_set = set(selected_categories)
        allow_null = "" in sel_set or "(Blank / Unspecified)" in sel_set

        str_col = df_sorted[split_column].cast(pl.Utf8)
        if allow_null:
            mask = str_col.is_in(list(sel_set)) | str_col.is_null() | (str_col.str.strip_chars() == "")
        else:
            mask = str_col.is_in(list(sel_set))
        df_target = df_sorted.filter(mask)
    else:
        df_target = df_sorted

    total_rows = df_target.height
    if total_rows == 0:
        raise ValueError("No matching records found for the selected categories.")

    # Mode: 'sorted_unified' (Single unified file sorted by country)
    if export_mode == "sorted_unified":
        ext = "xlsx" if file_format.lower() in ("xlsx", "excel") else "csv"
        safe_col = re.sub(r'[^a-zA-Z0-9_]', '_', split_column)
        out_name = f"{base_filename}_sorted_by_{safe_col}_{int(time.time())}.{ext}"
        out_path = os.path.join(output_dir, out_name)
        export_dataset(df_target, out_path, fmt=ext)
        return {
            "filename": out_name,
            "path": out_path,
            "mode": "sorted_unified",
            "total_rows": total_rows,
            "total_categories": df_target[split_column].n_unique(),
            "format": ext
        }

    # Partition by split_column
    partitions = df_target.partition_by([split_column], as_dict=True, maintain_order=True)

    # Mode: 'sheets' (Single Excel file with separate worksheets per Country!)
    if export_mode == "sheets":
        safe_col = re.sub(r'[^a-zA-Z0-9_]', '_', split_column)
        out_name = f"{base_filename}_by_{safe_col}_{int(time.time())}.xlsx"
        out_path = os.path.join(output_dir, out_name)

        wb = xlsxwriter.Workbook(out_path, {'constant_memory': False})
        header_format = wb.add_format({
            'bold': True,
            'font_color': '#FFFFFF',
            'bg_color': '#1E293B',
            'border': 1,
            'align': 'left'
        })
        summary_title_format = wb.add_format({
            'bold': True,
            'font_size': 13,
            'font_color': '#1E293B'
        })
        bold_format = wb.add_format({'bold': True})
        number_format = wb.add_format({'num_format': '#,##0'})
        pct_format = wb.add_format({'num_format': '0.0%'})

        seen_sheet_names = set()

        # Optional Overview / Summary sheet
        if include_summary_sheet:
            sum_ws = wb.add_worksheet("Overview_Summary")
            seen_sheet_names.add("overview_summary")
            sum_ws.freeze_panes(2, 0)
            sum_ws.write(0, 0, f"Summary: Grouped by '{split_column}'", summary_title_format)
            sum_ws.write_row(1, 0, [f"{split_column}", "Total Records", "Percentage of Dataset"], header_format)

            sum_row = 2
            sorted_parts = sorted(partitions.items(), key=lambda x: x[1].height, reverse=True)
            for part_key, part_df in sorted_parts:
                cat_val = part_key[0] if isinstance(part_key, tuple) else part_key
                cat_display = str(cat_val).strip() if cat_val is not None and str(cat_val).strip() else "(Blank / Unspecified)"
                cnt = part_df.height
                pct = (cnt / total_rows) if total_rows > 0 else 0
                sum_ws.write(sum_row, 0, cat_display)
                sum_ws.write(sum_row, 1, cnt, number_format)
                sum_ws.write(sum_row, 2, pct, pct_format)
                sum_row += 1

            # Total row
            sum_ws.write(sum_row, 0, "TOTAL", bold_format)
            sum_ws.write(sum_row, 1, total_rows, number_format)
            sum_ws.write(sum_row, 2, 1.0, pct_format)
            sum_ws.set_column(0, 0, 32)
            sum_ws.set_column(1, 1, 16)
            sum_ws.set_column(2, 2, 22)

        # Write each category as its own worksheet
        for part_key, part_df in partitions.items():
            cat_val = part_key[0] if isinstance(part_key, tuple) else part_key
            sheet_title = sanitize_sheet_title(str(cat_val) if cat_val is not None else "", seen_sheet_names)
            ws = wb.add_worksheet(sheet_title)
            ws.freeze_panes(1, 0)

            # Write header
            ws.write_row(0, 0, part_df.columns, header_format)

            # Write rows (handle max 1,048,500 rows per sheet)
            slice_rows = min(part_df.height, 1048500)
            for r_idx, row in enumerate(part_df.head(slice_rows).iter_rows()):
                ws.write_row(r_idx + 1, 0, row)

            if slice_rows > 0:
                ws.autofilter(0, 0, slice_rows, len(part_df.columns) - 1)

        wb.close()

        return {
            "filename": out_name,
            "path": out_path,
            "mode": "sheets",
            "total_rows": total_rows,
            "total_categories": len(partitions),
            "format": "xlsx"
        }

    # Mode: 'zip' (ZIP file with separate files per Country)
    elif export_mode == "zip":
        safe_col = re.sub(r'[^a-zA-Z0-9_]', '_', split_column)
        out_zip_name = f"{base_filename}_by_{safe_col}_{int(time.time())}.zip"
        out_zip_path = os.path.join(output_dir, out_zip_name)
        temp_split_dir = os.path.join(output_dir, f"_temp_{int(time.time())}_{uuid.uuid4().hex[:6]}")
        os.makedirs(temp_split_dir, exist_ok=True)

        try:
            file_ext = "xlsx" if file_format.lower() in ("xlsx", "excel") else "csv"
            seen_filenames = set()

            # Summary CSV
            summary_csv_path = os.path.join(temp_split_dir, "00_Overview_Summary.csv")
            summary_data = []

            for part_key, part_df in partitions.items():
                cat_val = part_key[0] if isinstance(part_key, tuple) else part_key
                raw_name = str(cat_val).strip() if cat_val is not None and str(cat_val).strip() else "Blank_Unspecified"
                safe_name = re.sub(r'[^a-zA-Z0-9_\- ]', '_', raw_name).strip()[:50]
                if not safe_name:
                    safe_name = "Category"
                cand_name = safe_name
                counter = 1
                while cand_name.lower() in seen_filenames:
                    cand_name = f"{safe_name}_{counter}"
                    counter += 1
                seen_filenames.add(cand_name.lower())

                part_file_name = f"{cand_name}.{file_ext}"
                part_file_path = os.path.join(temp_split_dir, part_file_name)

                if file_ext == "xlsx":
                    part_df.write_excel(part_file_path, worksheet="Cleaned_Data")
                else:
                    part_df.write_csv(part_file_path)

                summary_data.append({
                    "Category": raw_name,
                    "File": part_file_name,
                    "Rows": part_df.height
                })

            # Write summary CSV
            if summary_data:
                pl.DataFrame(summary_data).write_csv(summary_csv_path)

            # Zip the directory
            with zipfile.ZipFile(out_zip_path, 'w', compression=zipfile.ZIP_DEFLATED) as zipf:
                for root, _, files in os.walk(temp_split_dir):
                    for f in files:
                        full_f = os.path.join(root, f)
                        rel_f = os.path.relpath(full_f, temp_split_dir)
                        zipf.write(full_f, arcname=rel_f)

            return {
                "filename": out_zip_name,
                "path": out_zip_path,
                "mode": "zip",
                "total_rows": total_rows,
                "total_categories": len(partitions),
                "format": "zip"
            }
        finally:
            if os.path.exists(temp_split_dir):
                shutil.rmtree(temp_split_dir, ignore_errors=True)

    else:
        raise ValueError(f"Unsupported export mode '{export_mode}'.")

